# -*- coding: utf-8 -*-
"""pptxfig —— 生成「全可编辑」PPTX 插图的工具模块。

设计要点（踩过的坑，勿再犯）：
1. 【致命】OOXML 的 EMU 坐标必须是整数。Python 的 `/` 会产生 float，
   python-pptx 不做取整直接写入 XML（如 x="2414811.5"），PowerPoint/WPS
   解析失败后线条位置回退默认值 —— 表现就是箭头"飘"到别处。
   本模块所有坐标统一经 E() 取整。
2. 中文字体必须同时设置 latin 与 eastAsia 的 typeface，否则中文回落宋体。
3. a:ln 的子元素有严格顺序：fill → prstDash → headEnd → tailEnd，
   因此箭头端点必须在设置完颜色/虚线之后再 append。
4. 连接线方向由 flipH/flipV 表达：终点在起点左侧→flipH=1，上方→flipV=1。
"""
import re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

EMU_PER_INCH = 914400
EMU_PER_PT = 12700
EMU_PER_CM = 360000
EMU_PER_MM = 36000


def E(v):
    """强制取整为合法 EMU 整数（核心防御）。"""
    return int(round(v))


def inch(v):
    return E(v * EMU_PER_INCH)


def pt(v):
    return E(v * EMU_PER_PT)


def mm(v):
    return E(v * EMU_PER_MM)


def rgb(hex_str):
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_font(run, size_pt, bold=False, color=None, name="微软雅黑", italic=False):
    f = run.font
    f.name = name
    f.size = Pt(size_pt)
    f.bold = bold
    f.italic = italic
    if color is not None:
        f.color.rgb = color if isinstance(color, RGBColor) else rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


class Fig:
    """一页可编辑插图。所有坐标单位：英寸（内部转整数 EMU）。"""

    def __init__(self, width_in=13.333, height_in=7.5, bg=None):
        self.prs = Presentation()
        self.prs.slide_width = inch(width_in)
        self.prs.slide_height = inch(height_in)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self._shapes = []
        self._connectors = []
        if bg:
            sh = self.slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, self.prs.slide_height
            )
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(bg)
            sh.line.fill.background()
            sh.shadow.inherit = False

    # ---------- 基础图元 ----------
    def text(self, x, y, w, h, content, size=12, bold=False, color="#333333",
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, font="微软雅黑"):
        tb = self.slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        lines = content.split("\n")
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ln
            _set_font(r, size, bold, color, font, italic)
        return tb

    def box(self, x, y, w, h, title=None, desc=None, fill="#DEEBF7", line="#2E75B6",
            title_size=16, desc_size=11, title_color="#1F4E79", desc_color="#333333",
            line_w=1.5, round_adj=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, font="微软雅黑"):
        shp = self.slide.shapes.add_shape(shape, inch(x), inch(y), inch(w), inch(h))
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            try:
                shp.adjustments[0] = round_adj
            except Exception:
                pass
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(line_w)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = inch(0.06)
        tf.margin_top = tf.margin_bottom = inch(0.04)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        first = True
        if title:
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = title
            _set_font(r, title_size, True, title_color, font)
            first = False
        if desc:
            for ln in desc.split("\n"):
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                p.alignment = PP_ALIGN.CENTER
                p.space_before = Pt(3)
                r = p.add_run()
                r.text = ln
                _set_font(r, desc_size, False, desc_color, font)
        rec = {"left": inch(x), "top": inch(y), "w": inch(w), "h": inch(h),
               "name": title or desc or ""}
        self._shapes.append(rec)
        return shp

    def arrow(self, x1, y1, x2, y2, color="#C0392B", width=2.0, dash=None,
              head=False, tail=True, head_size="med"):
        """直线箭头。坐标单位：英寸。head=起点箭头，tail=终点箭头。"""
        X1, Y1, X2, Y2 = inch(x1), inch(y1), inch(x2), inch(y2)
        if X1 == X2 and Y1 == Y2:
            raise ValueError("箭头起止点重合")
        conn = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X1, Y1, X2, Y2)
        line = conn.line                      # LineFormat：用于颜色/宽度/虚线
        line.color.rgb = rgb(color)
        line.width = Pt(width)
        if dash:
            line.dash_style = dash
        ln = line._get_or_add_ln()            # 底层 XML：用于附加箭头端点
        # 注意：必须在 fill/dash 之后再 append，保证 a:ln 子元素顺序合法
        if head:
            el = ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": head_size, "len": head_size})
            ln.append(el)
        if tail:
            el = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": head_size, "len": head_size})
            ln.append(el)
        self._connectors.append({"x1": X1, "y1": Y1, "x2": X2, "y2": Y2,
                                 "color": color, "from": None, "to": None})
        return conn

    def arrow_between(self, box_a, box_b, side="h", color="#C0392B", width=2.0,
                      gap=0.05, dash=None, head=False, tail=True, offset=0.0):
        """在图框之间画箭头，自动贴合边缘（推荐用这个，避免算错端点）。
        side: 'h' 横向（a右→b左）；'v' 纵向（a下→b上）；'v-up' 纵向（a上→b下）。
        offset: 横向时相对垂直中心的偏移（英寸）。
        """
        a, b = box_a, box_b
        if side == "h":
            y = (a["top"] + a["h"] / 2.0) + offset
            x1 = a["left"] + a["w"] + inch(gap)
            x2 = b["left"] - inch(gap)
            self.arrow(x1 / EMU_PER_INCH, y / EMU_PER_INCH, x2 / EMU_PER_INCH, y / EMU_PER_INCH,
                       color, width, dash, head, tail)
        elif side == "v":
            x = a["left"] + a["w"] / 2.0
            y1 = a["top"] + a["h"] + inch(gap)
            y2 = b["top"] - inch(gap)
            self.arrow(x / EMU_PER_INCH, y1 / EMU_PER_INCH, x / EMU_PER_INCH, y2 / EMU_PER_INCH,
                       color, width, dash, head, tail)
        elif side == "v-up":
            x = a["left"] + a["w"] / 2.0
            y1 = a["top"] - inch(gap)
            y2 = b["top"] + b["h"] + inch(gap)
            self.arrow(x / EMU_PER_INCH, y1 / EMU_PER_INCH, x / EMU_PER_INCH, y2 / EMU_PER_INCH,
                       color, width, dash, head, tail)
        else:
            raise ValueError(side)
        self._connectors[-1]["from"] = a.get("name")
        self._connectors[-1]["to"] = b.get("name")

    def elbows(self, points, color="#2E75B6", width=1.5, dash=None,
               head=False, tail=True):
        """折线箭头（直角折线），points 为 (x,y) 英寸序列。仅末段带箭头。"""
        pts = [(inch(x), inch(y)) for x, y in points]
        for i in range(len(pts) - 1):
            last = (i == len(pts) - 2)
            conn = self.slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, pts[i][0], pts[i][1], pts[i + 1][0], pts[i + 1][1]
            )
            line = conn.line
            line.color.rgb = rgb(color)
            line.width = Pt(width)
            if dash:
                line.dash_style = dash
            ln = line._get_or_add_ln()
            if last and tail:
                ln.append(ln.makeelement(
                    qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
            if i == 0 and head:
                ln.append(ln.makeelement(
                    qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
            self._connectors.append({"x1": pts[i][0], "y1": pts[i][1],
                                     "x2": pts[i + 1][0], "y2": pts[i + 1][1],
                                     "color": color, "from": None, "to": None})
        return None

    def bar(self, x, y, w, h, label=None, fill="#2E5C8A", line=None,
            label_size=9, label_color="#FFFFFF", font="微软雅黑"):
        return self.box(x, y, w, h, title=label, fill=fill, line=line or fill,
                        title_size=label_size, title_color=label_color,
                        line_w=0.75, round_adj=0.25, font=font)

    # ---------- 输出 ----------
    def save(self, path):
        self.prs.save(path)
        return path


# ---------------- 自检 ----------------
def verify(pptx_path, tol_mm=2.5, verbose=True):
    """几何自检：
    1) XML 中所有 EMU 坐标必须为整数（浮点会导致位置丢失/箭头漂移）
    2) 所有图元必须在画布内
    3) 每条箭头必须指向（或起自）某个图框边缘 —— 否则就是"飘"了
    """
    issues = []
    prs = Presentation(pptx_path)
    sw, sh = prs.slide_width, prs.slide_height
    tol = mm(tol_mm)

    for si, slide in enumerate(prs.slides, 1):
        xml = slide.shapes._spTree.xml if hasattr(slide.shapes, "_spTree") else ""
        # 1) 整数坐标
        bad = re.findall(r'<a:(?:off|ext)[^>]*?="(-?\d+\.\d+)"', xml)
        if bad:
            issues.append(f"页{si}: 存在浮点 EMU 坐标 {len(bad)} 处（例 {bad[0]}）→ 位置会丢失")

        boxes, arrows = [], []
        for sp in slide.shapes:
            try:
                L, T = sp.left, sp.top
                W, H = sp.width, sp.height
            except (TypeError, AttributeError):
                continue
            if L is None:
                continue
            # 2) 画布范围
            if L < -tol or T < -tol or L + W > sw + tol or T + H > sh + tol:
                issues.append(f"页{si}: 图元超出画布 ({L},{T},{W},{H})")
            if sp.shape_type is not None and "AUTO_SHAPE" in str(sp.shape_type):
                txt = sp.text_frame.text.split("\n")[0] if sp.has_text_frame else ""
                boxes.append({"l": L, "t": T, "w": W, "h": H, "name": txt})
            elif "LINE" in str(sp.shape_type) or sp.__class__.__name__ == "Connector":
                # 3) 计算真实端点（考虑 flip）
                xfrm = sp._element.find(".//" + qn("a:xfrm"))
                off = xfrm.find(qn("a:off"))
                ext = xfrm.find(qn("a:ext"))
                ox, oy = int(off.get("x")), int(off.get("y"))
                cx, cy = int(ext.get("cx")), int(ext.get("cy"))
                fh = xfrm.get("flipH") in ("1", "true")
                fv = xfrm.get("flipV") in ("1", "true")
                sx = ox + cx if fh else ox
                sy = oy + cy if fv else oy
                ex = ox if fh else ox + cx
                ey = oy if fv else oy + cy
                ln = sp._element.find(".//" + qn("a:ln"))
                has_tail = ln is not None and ln.find(qn("a:tailEnd")) is not None
                has_head = ln is not None and ln.find(qn("a:headEnd")) is not None
                arrows.append({"s": (sx, sy), "e": (ex, ey),
                               "tail": has_tail, "head": has_head})

        n_checked = 0
        for ai, a in enumerate(arrows):
            # 仅校验「带箭头的那一端」：无箭头的参考线/折线中间段不参与
            targets = []
            if a.get("tail"):
                targets.append(("箭头端(终点)", a["e"]))
            if a.get("head"):
                targets.append(("箭头端(起点)", a["s"]))
            for tag, (px, py) in targets:
                n_checked += 1
                ok = False
                for b in boxes:
                    near_x = b["l"] - tol <= px <= b["l"] + b["w"] + tol
                    near_y = b["t"] - tol <= py <= b["t"] + b["h"] + tol
                    # 点落在框的"边缘邻域"内：横向邻域(贴左右边) 或 纵向邻域(贴上下边)
                    edge_x = (abs(px - b["l"]) <= tol) or (abs(px - (b["l"] + b["w"])) <= tol)
                    edge_y = (abs(py - b["t"]) <= tol) or (abs(py - (b["t"] + b["h"])) <= tol)
                    if (near_x and edge_y) or (near_y and edge_x):
                        ok = True
                        break
                    # 端点也可能落在框内部（如中心连线）
                    if near_x and near_y:
                        ok = True
                        break
                if not ok:
                    issues.append(
                        f"页{si}: 箭头#{ai + 1} 的{tag} ({px},{py}) 未贴近任何图框 → 疑似漂移")

        if verbose:
            print(f"  页{si}: 图框 {len(boxes)} 个 / 线条 {len(arrows)} 条 "
                  f"（其中箭头端 {n_checked} 个已校验贴合）")

    return issues


def export_png(pptx_path, out_dir, retries=3):
    """用本机 PowerPoint 导出每页 PNG（用于像素级校验）。返回 PNG 路径列表。

    COM 导出偶发返回空目录（PowerPoint 启动竞争），故重试；
    仍为空则显式报错，避免"静默失败"让校验形同虚设。
    """
    import os
    import time
    import win32com.client
    os.makedirs(out_dir, exist_ok=True)
    src = os.path.abspath(pptx_path)
    pngs = []
    for attempt in range(retries):
        app = win32com.client.Dispatch("PowerPoint.Application")
        try:
            pres = app.Presentations.Open(src, WithWindow=False)
            pres.SaveCopyAs(os.path.abspath(out_dir), 18)  # ppSaveAsPNG
            pres.Close()
        finally:
            app.Quit()
        pngs = sorted(
            os.path.join(out_dir, f) for f in os.listdir(out_dir)
            if f.lower().endswith(".png")
        )
        if pngs:
            return pngs
        time.sleep(1.0)  # 让上一次 PowerPoint 进程完全退出
    raise RuntimeError(
        f"PowerPoint 渲染失败：{retries} 次尝试均未导出 PNG（{out_dir}）。"
        f"请确认已安装 PowerPoint 且文件未被占用。")
