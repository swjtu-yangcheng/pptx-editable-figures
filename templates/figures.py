# -*- coding: utf-8 -*-
"""可编辑 PPTX 图型模板库（参数化，改数据即可复用）。

三种图型：
  make_flowchart()     横向流程图 + 闭环回流
  make_system_diagram() 三层体系图（上层汇聚→中心→下层辐射 + 回馈闭环）
  make_gantt()        甘特图（任务条 + 时间刻度 + 里程碑）

所有函数返回 Fig 对象；多个 Fig 可用 merge_figures() 合并为一个多页 PPTX。
"""
import copy
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))) + r"\scripts")
from pptx.enum.text import PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptxfig import Fig, EMU_PER_INCH

BLUE_D, BLUE_M, RED, GRAY = "#1F4E79", "#2E75B6", "#C0392B", "#595959"
W, H = 13.333, 7.5


# ---------------------------------------------------------------- 流程图
def make_flowchart(steps, title="流程图", loop_from=None, loop_to=None,
                   loop_label="", note="", size=None):
    """横向流程图。
    steps: [(标题, 描述), ...]，描述用 \n 换行
    loop_from/loop_to: 闭环回流的 源/目标 序号（0基），如 4 → 1
    """
    f = Fig(*(size or (W, H)))
    w, h = size or (W, H)
    f.text(0.5, 0.35, w - 1.0, 0.7, title, size=26, bold=True,
           color="#1F3F5C", align=PP_ALIGN.CENTER)

    n = len(steps)
    bw, bh, gap = 2.15, 1.90, 0.42
    total = bw * n + gap * (n - 1)
    x0 = (w - total) / 2.0
    top = 2.00
    for i, (t, d) in enumerate(steps):
        f.box(x0 + i * (bw + gap), top, bw, bh, title=t, desc=d,
              fill="#E8F1F8", line=BLUE_M, title_size=17, desc_size=11,
              title_color=BLUE_D, desc_color="#33506B")
    recs = [{"left": (x0 + i * (bw + gap)) * EMU_PER_INCH, "top": top * EMU_PER_INCH,
             "w": bw * EMU_PER_INCH, "h": bh * EMU_PER_INCH, "name": s[0]}
            for i, s in enumerate(steps)]
    for i in range(n - 1):
        f.arrow_between(recs[i], recs[i + 1], "h", color=RED, width=2.5, gap=0.03)

    if loop_from is not None and loop_to is not None:
        y_loop = top + bh + 0.55
        xs = (x0 + loop_from * (bw + gap) + bw / 2.0)
        xe = (x0 + loop_to * (bw + gap) + bw / 2.0)
        f.elbows([(xs, top + bh), (xs, y_loop), (xe, y_loop), (xe, top + bh + 0.03)],
                 color=BLUE_M, width=1.5, dash=MSO_LINE_DASH_STYLE.DASH, tail=True)
        if loop_label:
            f.text(xe + 0.35, y_loop - 0.30, 4.8, 0.35, loop_label, size=12, color=GRAY)
    if note:
        f.text(1.0, h - 1.05, w - 2.0, 0.45, note, size=13, color=GRAY,
               align=PP_ALIGN.CENTER)
    return f


# ---------------------------------------------------------------- 体系图
def make_system_diagram(upper, center, lower, title="体系图",
                        loop_pair=(0, 0), loop_label="回馈反哺", note="", size=None):
    """三层体系图：上层(汇聚)→中心→下层(辐射)，左侧可加回馈闭环。
    upper/lower: [(标题, 描述, 填充色), ...]，各 3 项效果最佳
    center: 中心框文字
    loop_pair: (下层序号, 上层序号) 用于画回馈闭环
    """
    w, h = size or (W, H)
    f = Fig(w, h)
    f.text(0.5, 0.35, w - 1.0, 0.7, title, size=26, bold=True,
           color="#1F3F5C", align=PP_ALIGN.CENTER)

    CW = 2.6
    COL = [1.35, 5.55, 9.75]
    Y_UP, H_UP = 1.55, 1.15
    Y_CT, H_CT = 3.25, 1.05
    Y_DN, H_DN = 4.85, 1.15
    cxs = [COL[1] + 0.7, COL[1] + CW / 2.0, COL[1] + CW - 0.7]

    for i, (t, d, c) in enumerate(upper):
        f.box(COL[i], Y_UP, CW, H_UP, title=t, desc=d, fill=c, line=BLUE_M,
              title_size=15, desc_size=10.5, title_color=BLUE_D, desc_color="#33506B")
        f.arrow(COL[i] + CW / 2.0, Y_UP + H_UP + 0.02, cxs[i], Y_CT - 0.02,
                color=BLUE_M, width=1.8)
    for i, (t, d, c) in enumerate(lower):
        f.box(COL[i], Y_DN, CW, H_DN, title=t, desc=d, fill=c, line=BLUE_M,
              title_size=15, desc_size=10.5, title_color=BLUE_D, desc_color="#33506B")
        f.arrow(cxs[i], Y_CT + H_CT + 0.02, COL[i] + CW / 2.0, Y_DN - 0.02,
                color=BLUE_M, width=1.8)
    f.box(COL[1], Y_CT, CW, H_CT, title=center, fill="#2E5C8A", line="#FFFFFF",
          title_size=15, title_color="#FFFFFF", line_w=1.5)

    if loop_pair:
        li, ui = loop_pair
        y_from, y_to = Y_DN + H_DN / 2.0, Y_UP + H_UP / 2.0
        f.elbows([(COL[li] - 0.02, y_from), (0.85, y_from), (0.85, y_to),
                  (COL[ui] - 0.02, y_to)],
                 color=RED, width=1.6, dash=MSO_LINE_DASH_STYLE.DASH, tail=True)
        if loop_label:
            f.text(0.20, 3.05, 0.5, 1.5, "\n".join(loop_label), size=12,
                   bold=True, color=RED, align=PP_ALIGN.CENTER)
    if note:
        f.text(0.5, 6.35, w - 1.0, 0.4, note, size=11.5, color=GRAY,
               align=PP_ALIGN.CENTER)
    return f


# ---------------------------------------------------------------- 甘特图
def make_gantt(tasks, milestones, start_year=2026, start_month=9, months=22,
               title="甘特图", note="", size=None):
    """甘特图。
    tasks: [(任务名, 起始月序, 持续月数, 颜色), ...]
    milestones: {月序: 名称}
    """
    w, h = size or (W, H)
    f = Fig(w, h)
    f.text(0.5, 0.30, w - 1.0, 0.62, title, size=24, bold=True,
           color="#1F3F5C", align=PP_ALIGN.CENTER)

    X0, X1 = 3.30, 12.55
    WPM = (X1 - X0) / float(months)
    mx = lambda m: X0 + m * WPM
    ROW0, ROWH, BARH, AXIS_Y = 1.95, 0.66, 0.40, 5.02

    for m in range(0, months + 1, 3):
        f.arrow(mx(m), 1.72, mx(m), AXIS_Y, color="#D9D9D9", width=0.75, tail=False)
    f.arrow(X0, AXIS_Y, X1, AXIS_Y, color="#7F7F7F", width=1.25, tail=False)

    for i, (name, start, dur, color) in enumerate(tasks):
        # 显式校验：第3个参数最容易被误当成"结束月序"，静默超界比报错更难查
        if start < 0 or start + dur > months:
            raise ValueError(
                f"任务「{name}」区间 {start}~{start + dur} 超出时间轴 0~{months} 个月。"
                f"注意第 3 个参数是【持续月数】，不是【结束月序】；"
                f"若想表示 {start} 到 {dur}，应写 ({start}, {dur - start})。")
        y = ROW0 + i * ROWH
        f.text(0.45, y - 0.02, 2.75, 0.42, name, size=12.5, color="#333333")
        f.bar(mx(start), y, dur * WPM, BARH, label=f"{dur}个月", fill=color,
              line=color, label_size=10.5, label_color="#FFFFFF")

    for m in range(0, months + 1, 3):
        m0 = start_month + m
        yr = start_year + (m0 - 1) // 12
        mo = (m0 - 1) % 12 + 1
        f.text(mx(m) - 0.55, AXIS_Y + 0.08, 1.1, 0.32, f"{yr}.{mo:02d}",
               size=11, color="#404040", align=PP_ALIGN.CENTER)

    for m, txt in sorted(milestones.items()):
        f.arrow(mx(m), 1.68, mx(m), AXIS_Y, color=RED, width=1.0,
                dash=MSO_LINE_DASH_STYLE.DASH, tail=False)
        f.text(mx(m) - 0.62, 1.28, 1.24, 0.32, txt, size=11, bold=True,
               color=RED, align=PP_ALIGN.CENTER)

    f.text(0.45, 6.15, 6.0, 0.4, "■ 条形：阶段工作周期", size=12, color="#2E5C8A")
    f.text(6.4, 6.15, 6.5, 0.4, "--- 红色虚线：里程碑节点", size=12, color=RED)
    if note:
        f.text(0.45, 6.62, w - 0.9, 0.4, note, size=11.5, color=GRAY)
    return f


# ---------------------------------------------------------------- 合并
def merge_figures(figs, out_path):
    """把多个 Fig 合并为一个多页 PPTX。注意：会就地修改 figs[0]。"""
    base = figs[0].prs
    for extra in figs[1:]:
        dst = base.slides.add_slide(base.slide_layouts[6])
        for shp in list(extra.slide.shapes):
            dst.shapes._spTree.append(copy.deepcopy(shp._element))
    base.save(out_path)
    return out_path


# ---------------------------------------------------------------- 示例
if __name__ == "__main__":
    import tempfile
    demo = os.path.join(tempfile.gettempdir(), "pptxfig_demo.pptx")
    f1 = make_flowchart(
        [("信息采集", "学业数据、行为数据\n建立学生学业档案"),
         ("分级预警", "黄/橙/红三级\n教务初筛、辅导员复核"),
         ("预警谈话", "48小时内约谈\n共定改进计划"),
         ("多方联动", "家长告知、教师联动\n组建帮扶小组"),
         ("跟踪销号", "月度跟进、定期会商\n达标解除预警")],
        title="图2-1  学业预警与帮扶工作流程图", loop_from=4, loop_to=1,
        loop_label="未改善：升级干预措施，必要时转介专业力量",
        note="闭环管理：预警不解除不销号，销号后继续跟踪一个学期")

    f2 = make_system_diagram(
        [("物质帮助", "奖助勤贷免补\n保障基本就学", "#D6E6F2"),
         ("能力拓展", "勤工助学、技能培训\n访学交流、就业支持", "#BBD6EA"),
         ("精神激励", "励志教育\n感恩诚信教育", "#D6E6F2")],
        "解困—育人—成才—回馈",
        [("价值引领", "志愿服务\n资助宣传大使", "#BBD6EA"),
         ("精准认定", "定量+定性\n动态调整", "#D6E6F2"),
         ("隐私保护", "去标识化公示\n限定数据用途", "#BBD6EA")],
        title="图8-1  发展型资助育人体系构成", loop_pair=(0, 0), loop_label="回馈反哺",
        note="说明：上层为资助途径，中层为育人目标链条，下层为支撑保障机制。")

    f3 = make_gantt(
        [("一  大纲论证与体例统一", 0, 3, "#2E5C8A"),
         ("二  分篇撰写初稿", 3, 7, "#3E7CB1"),
         ("三  教学试用与反馈", 10, 3, "#4FA3D1"),
         ("四  统稿修改与专家鉴定", 13, 3, "#6FB7DC"),
         ("五  定稿出版与推广", 16, 6, "#8FC7E8")],
        {3: "初稿完成", 10: "试用启动", 13: "修改启动", 16: "提交出版", 22: "出版发行"},
        title="图3-1  教材编写工作进度甘特图（2026.09—2028.06）",
        note="注：编写周期共22个月，第五阶段结束于2028年6月。")

    merge_figures([f1, f2, f3], demo)
    print("demo saved:", demo)
    from pptxfig import verify
    print(verify(demo) or "自检通过")
